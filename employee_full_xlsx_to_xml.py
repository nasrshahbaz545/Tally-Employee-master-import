#!/usr/bin/env python3
from openpyxl import load_workbook
from pathlib import Path
from xml.sax.saxutils import escape, quoteattr
import argparse, re, sys

HEADERS=['EmployeeName','EmployeeCategory','EmployeeGroup','EmployeeNumber','DateOfJoining','EmploymentStatus','Designation','Function','Location','Gender','DateOfBirth','BloodGroup','FatherMotherName','SpouseName','AddressLine1','AddressLine2','ContactNumber','Email','BankName','BankBranch','BankAccountNumber','MICRCode','CNIC','CountryOfIssue','PassportExpiryDate','VisaNumber','VisaExpiryDate','WorkPermitNumber','ContractStartDate','ContractExpiryDate','DateOfLeaving','ReasonForLeaving','PANNumber','PFAccountNumber','FPFAccountNumber','PFJoiningDate','PFRelievingDate','ESINumber','ESIDispensaryName']
DATE_FIELDS={'DateOfJoining':'DATEOFJOIN','DateOfBirth':'DATEOFBIRTH','PassportExpiryDate':'PASSPORTEXPIRYDATE','VisaExpiryDate':'VISAEXPIRYDATE','ContractStartDate':'CONTRACTSTARTDATE','ContractExpiryDate':'CONTRACTEXPIRYDATE','DateOfLeaving':'DEACTIVATIONDATE','PFJoiningDate':'PFJOININGDATE','PFRelievingDate':'PFRELIEVINGDATE'}
TAG_FIELDS={'Designation':'DESIGNATION','Function':'FUNCTION','Location':'LOCATION','Gender':'GENDER','BloodGroup':'BLOODGROUP','FatherMotherName':'FATHERNAME','SpouseName':'SPOUSENAME','ContactNumber':'CONTACTNUMBERS','Email':'EMAILID','BankName':'BANKDETAILS','BankBranch':'BANKBRANCH','BankAccountNumber':'BANKACCOUNTNUMBER','MICRCode':'MICRCODE','CNIC':'PASSPORTDETAILS','CountryOfIssue':'COUNTRYOFISSUE','VisaNumber':'VISANUMBER','WorkPermitNumber':'WORKPERMITNUMBER','ReasonForLeaving':'REASONSFORLEAVING','PANNumber':'PANNUMBER','PFAccountNumber':'PFACCOUNTNUMBER','FPFAccountNumber':'FPFACCOUNTNUMBER','ESINumber':'ESINUMBER','ESIDispensaryName':'ESIDISPENSARYNAME'}
STATUSES={'Regular','Contract','Deputation','Daily Wages','Intern','Project Employee','Retired','Terminated'}

def st(v): return '' if v is None else str(v).strip()
def xe(v): return escape(st(v))
def dt(v,row,name):
    if v is None or st(v)=='': return ''
    if hasattr(v,'strftime'): return v.strftime('%Y%m%d')
    q=re.sub(r'\D','',st(v))
    if len(q)!=8: raise ValueError(f'Row {row}: {name} must be YYYYMMDD')
    return q

def read(path):
    wb=load_workbook(path,data_only=True); ws=wb['Employees']; hh=[st(c.value) for c in ws[1]]
    miss=[h for h in HEADERS if h not in hh]
    if miss: raise ValueError('Missing columns: '+', '.join(miss))
    c={h:hh.index(h)+1 for h in HEADERS}; data=[]; names=set(); nums=set(); cnics=set()
    for r in range(2,ws.max_row+1):
        d={h:st(ws.cell(r,c[h]).value) for h in HEADERS}
        if not d['EmployeeName']: continue
        if not d['EmployeeCategory'] or not d['EmployeeGroup']: raise ValueError(f'Row {r}: EmployeeCategory and EmployeeGroup are required')
        if d['EmploymentStatus'] and d['EmploymentStatus'] not in STATUSES: raise ValueError(f"Row {r}: invalid EmploymentStatus '{d['EmploymentStatus']}'")
        for h in DATE_FIELDS: d[h]=dt(ws.cell(r,c[h]).value,r,h)
        nk=d['EmployeeName'].casefold()
        if nk in names: raise ValueError(f"Row {r}: duplicate EmployeeName '{d['EmployeeName']}'")
        names.add(nk)
        if d['EmployeeNumber']:
            k=d['EmployeeNumber'].casefold()
            if k in nums: raise ValueError(f"Row {r}: duplicate EmployeeNumber '{d['EmployeeNumber']}'")
            nums.add(k)
        if d['CNIC']:
            if d['CNIC'] in cnics: raise ValueError(f"Row {r}: duplicate CNIC '{d['CNIC']}'")
            cnics.add(d['CNIC'])
        data.append(d)
    if not data: raise ValueError('No employee rows found')
    return data

def tag(n,v,ind='      '): return f'{ind}<{n}>{xe(v)}</{n}>\n' if st(v) else ''
def block(d):
    b='    <TALLYMESSAGE xmlns:UDF="TallyUDF">\n'
    b+=f'     <COSTCENTRE NAME={quoteattr(d["EmployeeName"])} RESERVEDNAME="">\n'
    if d['EmployeeNumber']:
        b+='      <MAILINGNAME.LIST TYPE="String">\n'+tag('MAILINGNAME',d['EmployeeNumber'],'       ')+'      </MAILINGNAME.LIST>\n'
    b+=tag('DATEOFJOIN',d['DateOfJoining'])+tag('DEACTIVATIONDATE',d['DateOfLeaving'])
    b+=tag('PARENT',d['EmployeeGroup'])+tag('CATEGORY',d['EmployeeCategory'])
    for h,t in TAG_FIELDS.items(): b+=tag(t,d[h])
    for h,t in DATE_FIELDS.items():
        if h not in ('DateOfJoining','DateOfLeaving'): b+=tag(t,d[h])
    addresses=[d['AddressLine1'],d['AddressLine2']]
    if any(addresses):
        b+='      <ADDRESS.LIST TYPE="String">\n'
        for a in addresses: b+=tag('ADDRESS',a,'       ')
        b+='      </ADDRESS.LIST>\n'
    b+='      <AFFECTSSTOCK>No</AFFECTSSTOCK>\n      <FORPAYROLL>Yes</FORPAYROLL>\n      <FORJOBCOSTING>No</FORJOBCOSTING>\n      <ISEMPLOYEEGROUP>No</ISEMPLOYEEGROUP>\n      <SORTPOSITION> 1000</SORTPOSITION>\n'
    b+='      <LANGUAGENAME.LIST>\n       <NAME.LIST TYPE="String">\n'+tag('NAME',d['EmployeeName'],'        ')+'       </NAME.LIST>\n       <LANGUAGEID> 1033</LANGUAGEID>\n      </LANGUAGENAME.LIST>\n'
    if d['EmploymentStatus']:
        b+='      <UDF:NCRCEMPLOYMENTSTATUS.LIST DESC="`NCRCEmploymentStatus`" ISLIST="YES" TYPE="String">\n'
        b+=f'       <UDF:NCRCEMPLOYMENTSTATUS DESC="`NCRCEmploymentStatus`">{xe(d["EmploymentStatus"])}</UDF:NCRCEMPLOYMENTSTATUS>\n'
        b+='      </UDF:NCRCEMPLOYMENTSTATUS.LIST>\n'
    b+='     </COSTCENTRE>\n    </TALLYMESSAGE>\n'; return b

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('input',nargs='?',default='Employee_Master_Full_Template.xlsx'); ap.add_argument('output',nargs='?',default='Employee_Master_Full_Import.xml'); ap.add_argument('--company',default='National Commission on the Rights of Child'); a=ap.parse_args()
    rows=read(a.input)
    z='<?xml version="1.0" encoding="UTF-8"?>\n<ENVELOPE>\n <HEADER>\n  <TALLYREQUEST>Import Data</TALLYREQUEST>\n </HEADER>\n <BODY>\n  <IMPORTDATA>\n   <REQUESTDESC>\n    <REPORTNAME>All Masters</REPORTNAME>\n    <STATICVARIABLES>\n     <SVCURRENTCOMPANY>'+xe(a.company)+'</SVCURRENTCOMPANY>\n    </STATICVARIABLES>\n   </REQUESTDESC>\n   <REQUESTDATA>\n'+''.join(block(r) for r in rows)+'   </REQUESTDATA>\n  </IMPORTDATA>\n </BODY>\n</ENVELOPE>\n'
    Path(a.output).write_text(z,encoding='utf-8'); print(f'Created {a.output} with {len(rows)} employee record(s).')
if __name__=='__main__':
    try: main()
    except Exception as e: print('ERROR:',e,file=sys.stderr); raise SystemExit(1)
